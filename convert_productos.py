#!/usr/bin/env python3
"""Convierte la hoja Productos a public/products.json sin publicar datos inválidos.

Uso: python3 convert_productos.py [willitboard_productos.xlsx] [public/products.json] [--report informe.txt]

Mismas reglas que los operadores (convert.py):
- Un producto se publica SOLO si tiene verificado_el puesto por una persona,
  fuente_url (página del FABRICANTE) y al menos un enlace de tienda.
- estado tiene que ser "listo" o "publicado"; "candidato" y "retirado" no salen.
- Bolsos y maletas necesitan las tres medidas exteriores: sin medidas no hay veredicto.
- NUNCA se publican precio_ref, notas ni las columnas de vista previa: son internas.
- Si hay errores en una fila lista para publicar, se conserva el JSON anterior.

Este archivo NO toca operators.json: el dataset de operadores sigue siendo uno solo.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

from convert import atomic_write, boolean, number, review_date, source_url, text

SCHEMA_VERSION = "1.0"
STALE_DAYS = 180
HEADER_ROW = 1
FIRST_DATA_ROW = 3          # la fila 2 son las ayudas de cada columna
CATEGORIES = ("underseat-bags", "cabin-backpacks", "cabin-suitcases", "packing", "liquids",
              "weigh-measure", "tech", "comfort", "security")
BAG_CATEGORIES = ("underseat-bags", "cabin-backpacks", "cabin-suitcases")
PUBLISHABLE = ("listo", "publicado")
STATES = ("candidato", "listo", "publicado", "retirado")
PRICE_BANDS = ("€", "€€", "€€€")
STORE_COLUMNS = {"link_amazon_de": "amazon_de", "link_amazon_se": "amazon_se",
                 "link_amazon_uk": "amazon_uk", "link_otro": "other"}
REQUIRED_HEADERS = ("id", "nombre", "marca", "categoria", "largo_cm", "ancho_cm", "alto_cm",
                    "medidas_incluyen_ruedas", "peso_vacio_kg", "capacidad_l", "rango_precio",
                    "fuente_url", "estado", "verificado_el") + tuple(STORE_COLUMNS)


def load_rows(src):
    from openpyxl import load_workbook
    wb = load_workbook(src, read_only=True, data_only=True)
    if "Productos" not in wb.sheetnames:
        raise ValueError("la planilla no tiene la hoja Productos")
    ws = wb["Productos"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [text(h) for h in rows[HEADER_ROW - 1]]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ValueError("faltan columnas: " + ", ".join(missing))
    out = []
    for n, values in enumerate(rows[FIRST_DATA_ROW - 1:], start=FIRST_DATA_ROW):
        row = dict(zip(headers, values))
        if text(row.get("id")) is None:
            continue
        out.append((n, row))
    return out


def build_product(row, line, today):
    """Devuelve (producto | None, motivo_si_no_se_publica, errores)."""
    pid = text(row["id"])
    estado = (text(row["estado"]) or "candidato").lower()
    errors = []
    if estado not in STATES:
        return None, None, [f"fila {line} ({pid}): estado '{estado}' no existe; usar " + ", ".join(STATES)]
    if estado not in PUBLISHABLE:
        return None, f"estado {estado}", []

    def read(field, parser, *args):
        try:
            return parser(row.get(field), *args)
        except ValueError as exc:
            errors.append(f"fila {line} ({pid}), {field}: {exc}")
            return None

    verified = read("verificado_el", review_date, today)
    if verified is None and not errors:
        return None, "falta verificado_el puesto por una persona", []
    src = read("fuente_url", source_url)
    cat = text(row["categoria"])
    if cat not in CATEGORIES:
        errors.append(f"fila {line} ({pid}), categoria: usar una de " + ", ".join(CATEGORIES))
    dims = [read(f, number) for f in ("largo_cm", "ancho_cm", "alto_cm")]
    have = [d for d in dims if d is not None]
    if 0 < len(have) < 3:
        errors.append(f"fila {line} ({pid}): medidas incompletas; cargar las tres o ninguna")
    if cat in BAG_CATEGORIES and len(have) != 3:
        errors.append(f"fila {line} ({pid}): un bolso o maleta necesita las tres medidas exteriores")
    links = {}
    for col, key in STORE_COLUMNS.items():
        if text(row.get(col)) is not None:
            url = read(col, source_url)
            if url:
                links[key] = url
    if not links:
        errors.append(f"fila {line} ({pid}): sin enlace de tienda no hay nada que publicar")
    band = text(row.get("rango_precio"))
    if band is not None and band not in PRICE_BANDS:
        errors.append(f"fila {line} ({pid}), rango_precio: usar €, €€ o €€€")
    wheels = read("medidas_incluyen_ruedas", boolean)
    kg = read("peso_vacio_kg", number)
    cap = read("capacidad_l", number)
    if errors:
        return None, None, errors
    return {
        "id": pid,
        "name": text(row["nombre"]),
        "brand": text(row["marca"]),
        "category": cat,
        "dims_cm": dims if len(have) == 3 else None,
        "dims_include_wheels": wheels,
        "empty_kg": kg,
        "capacity_l": cap,
        "price_band": band,
        "links": links,
        "source_url": src,
        "verified_on": verified.isoformat(),
        "days_since_verified": (today - verified).days,
    }, None, []


def convert_rows(rows, today):
    products, skipped, errors, seen = [], [], [], set()
    for line, row in rows:
        pid = text(row["id"])
        if pid in seen:
            errors.append(f"fila {line}: id repetido '{pid}'")
            continue
        seen.add(pid)
        product, reason, errs = build_product(row, line, today)
        errors += errs
        if product:
            products.append(product)
        elif reason:
            skipped.append(f"fila {line} ({pid}): no se publica — {reason}")
    return products, skipped, errors


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", nargs="?", default="willitboard_productos.xlsx")
    ap.add_argument("destination", nargs="?", default="public/products.json")
    ap.add_argument("--report", type=Path)
    args = ap.parse_args(argv)
    today = date.today()
    try:
        rows = load_rows(args.source)
    except (OSError, ValueError) as exc:
        print(f"No se pudo leer la planilla: {exc}", file=sys.stderr)
        return 2
    products, skipped, errors = convert_rows(rows, today)
    lines = skipped[:]
    stale = [p["id"] for p in products if p["days_since_verified"] > STALE_DAYS]
    if stale:
        lines.append("Revisión vencida (más de %d días): %s" % (STALE_DAYS, ", ".join(stale)))
    if errors:
        lines = ["ERRORES — se conserva el products.json anterior:"] + errors + lines
        code = 1
    else:
        data = {"schema_version": SCHEMA_VERSION, "generated_on": today.isoformat(),
                "canonical_units": {"length": "cm", "weight": "kg"}, "products": products}
        atomic_write(args.destination, json.dumps(data, ensure_ascii=False, indent=2) + "\n")
        lines.append(f"Publicados {len(products)} productos en {args.destination}.")
        code = 0
    report = "\n".join(lines)
    print(report)
    if args.report:
        args.report.write_text(report + "\n", encoding="utf-8")
    return code


if __name__ == "__main__":
    sys.exit(main())
