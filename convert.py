#!/usr/bin/env python3
"""Convierte la hoja Operadores a JSON 1.4 sin publicar datos inválidos.

Uso: python convert.py planilla.xlsx public/operators.json --report informe.txt
Requiere Python 3.9+ y openpyxl. La fecha de revisión la carga una persona.
"""

from __future__ import annotations

import argparse
import ipaddress
import json
import math
import os
import re
import sys
import tempfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import BadZipFile


SCHEMA_VERSION = "1.4"
STALE_DAYS = 180
EXAMPLE_MARKER = "EJEMPLO SIN VERIFICAR"
REQUIRED_HEADERS = (
    "id nombre tipo pais region limit_type personal_incluido "
    "personal_largo_cm personal_ancho_cm personal_alto_cm personal_kg "
    "cabina_incluida cabina_addon cabina_largo_cm cabina_ancho_cm cabina_alto_cm "
    "cabina_kg incluye_ruedas_manijas tolerancia_ruedas_manijas_cm "
    "moneda tarifa_online_desde tarifa_en_puerta_desde "
    "regla_texto_es regla_texto_en notas fuente_url verificado_el"
).split()
BOOLEAN_FIELDS = ("personal_incluido", "cabina_incluida", "incluye_ruedas_manijas")
POSITIVE_FIELDS = (
    "personal_largo_cm", "personal_ancho_cm", "personal_alto_cm", "personal_kg",
    "cabina_largo_cm", "cabina_ancho_cm", "cabina_alto_cm", "cabina_kg",
)
FEE_FIELDS = ("tarifa_online_desde", "tarifa_en_puerta_desde")
NONNEGATIVE_FIELDS = ("tolerancia_ruedas_manijas_cm",)
# Columna opcional desde el esquema 1.4 (24/09/2026): el nombre comercial de la
# tarifa a la que corresponden las cifras, tal como lo escribe el operador
# ("Basic Fare", "Standard"...). No se traduce. Opcional para que una planilla
# vieja siga convirtiendo: vacío significa "no confirmado", y la ficha muestra
# entonces el aviso genérico de tarifas.
OPTIONAL_HEADERS = ("tarifa_referencia", "peso_combinado_kg")
# peso_combinado_kg (esquema 1.4, 26/09/2026): algunos operadores publican UN
# peso para el artículo personal y la valija JUNTOS (Air France, KLM, Transavia,
# Norwegian). Se guarda aparte del peso por bulto: dejarlo vacío haría decir a
# la ficha "sin peso publicado", que sería falso, y repartirlo entre los dos
# bultos sería inventar una cifra que el operador nunca publicó.
MAX_FARE_LEN = 80
SLUG = re.compile(r"[a-z0-9]+(?:-[a-z0-9]+)*\Z")
DECIMAL = re.compile(r"[+-]?(?:[0-9]+(?:[.,][0-9]+)?|[.,][0-9]+)\Z")


def text(value):
    if value is None:
        return None
    return str(value).strip() or None


def boolean(value):
    if text(value) is None:
        return None
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().upper()
    if normalized in ("SI", "SÍ"):
        return True
    if normalized == "NO":
        return False
    raise ValueError("usar SI, NO o una celda vacía; un valor desconocido no significa NO")


def number(value, allow_zero=False):
    if text(value) is None:
        return None
    if isinstance(value, bool):
        raise ValueError("se esperaba un número, no un valor lógico")
    if isinstance(value, str):
        if not DECIMAL.fullmatch(value.strip()):
            raise ValueError("usar un número sin unidades ni separadores de miles")
        value = value.strip().replace(",", ".")
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        raise ValueError("se esperaba un número") from None
    if not math.isfinite(result):
        raise ValueError("el número debe ser finito")
    if result < 0 or (result == 0 and not allow_zero):
        raise ValueError("el número debe ser mayor o igual a cero" if allow_zero
                         else "el número debe ser mayor que cero")
    return int(result) if result.is_integer() else result


def review_date(value, today):
    if text(value) is None:
        return None
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    else:
        raw = str(value).strip()
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", raw):
            raise ValueError("usar una fecha de Excel o texto completo AAAA-MM-DD")
        try:
            result = date.fromisoformat(raw)
        except ValueError:
            raise ValueError("la fecha no existe") from None
    if result > today:
        raise ValueError("la fecha de revisión no puede estar en el futuro")
    return result


def source_url(value):
    raw = text(value)
    if raw is None:
        raise ValueError("falta la URL de la fuente; la fecha sola no permite publicar")
    try:
        parsed = urlsplit(raw)
        hostname = parsed.hostname
        parsed.port
        if parsed.scheme not in ("http", "https") or not hostname:
            raise ValueError
        if parsed.username or parsed.password or re.search(r"\s", raw):
            raise ValueError
        hostname = hostname.rstrip(".").encode("idna").decode("ascii")
        if len(hostname) > 253 or "." not in hostname:
            raise ValueError
        labels = hostname.split(".")
        if not all(re.fullmatch(r"[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?", label)
                   for label in labels):
            raise ValueError
        try:
            ipaddress.ip_address(hostname)
        except ValueError:
            pass
        else:
            raise ValueError
    except (ValueError, UnicodeError):
        raise ValueError("usar una URL http/https con nombre de dominio válido") from None
    return raw


def enum(value, choices, uppercase=False):
    raw = text(value)
    normalized = (raw.upper() if uppercase else raw.lower()) if raw else None
    if normalized not in choices:
        raise ValueError("usar uno de: " + ", ".join(choices))
    return normalized


def load_rows(src):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(src, data_only=False, read_only=True)
    except InvalidFileException as exc:
        raise ValueError(str(exc)) from None
    try:
        if "Operadores" not in workbook.sheetnames:
            raise ValueError("falta la hoja Operadores")
        sheet = workbook["Operadores"]
        headers = [text(cell.value) for cell in sheet[1]]
        repeated = sorted(h for h, count in Counter(h for h in headers if h).items() if count > 1)
        missing = sorted(set(REQUIRED_HEADERS) - set(headers))
        if repeated or missing:
            details = []
            if missing:
                details.append("faltan columnas: " + ", ".join(missing))
            if repeated:
                details.append("columnas duplicadas: " + ", ".join(repeated))
            raise ValueError("; ".join(details))

        rows, examples = [], []
        for line, cells in enumerate(sheet.iter_rows(min_row=3), start=3):
            row = {headers[i]: cell.value for i, cell in enumerate(cells)
                   if i < len(headers) and headers[i]}
            if all(text(value) is None for value in row.values()):
                continue
            if EXAMPLE_MARKER in (text(row.get("notas")) or "").upper():
                examples.append(line)
                continue
            rows.append((line, row))
        return rows, examples
    finally:
        workbook.close()


def build_operator(row, line, today):
    errors, warnings = [], []

    def read(field, parser):
        try:
            return parser(row.get(field))
        except ValueError as exc:
            errors.append(f"Fila {line} ({text(row.get('id')) or 'sin id'}), {field}: {exc}.")
            return None

    name = text(row.get("nombre"))
    if name is None:
        errors.append(f"Fila {line}: falta nombre.")
    operator_type = read("tipo", lambda v: enum(v, ("air", "rail", "ferry")))
    region = read("region", lambda v: enum(v, ("EU", "LATAM", "NA", "ASIA", "OTHER"), True))
    limit_type = read("limit_type", lambda v: enum(v, ("dimensional", "descriptivo", "ninguno")))
    country = (text(row.get("pais")) or "").upper()
    if not re.fullmatch(r"[A-Z]{2}", country):
        errors.append(f"Fila {line}, pais: usar un código de dos letras, por ejemplo ES.")
    verified = read("verificado_el", lambda v: review_date(v, today))
    source = read("fuente_url", source_url)
    booleans = {field: read(field, boolean) for field in BOOLEAN_FIELDS}
    numbers = {field: read(field, number) for field in POSITIVE_FIELDS}
    numbers.update({field: read(field, lambda v: number(v, allow_zero=True)) for field in FEE_FIELDS})
    numbers.update({field: read(field, lambda v: number(v, allow_zero=True)) for field in NONNEGATIVE_FIELDS})
    currency = text(row.get("moneda"))
    if currency:
        currency = currency.upper()
        if not re.fullmatch(r"[A-Z]{3}", currency):
            errors.append(f"Fila {line}, moneda: usar un código de tres letras, por ejemplo EUR.")
    if any(numbers[field] is not None for field in FEE_FIELDS) and not currency:
        errors.append(f"Fila {line}, moneda: es obligatoria cuando se carga una tarifa.")

    tolerance = numbers["tolerancia_ruedas_manijas_cm"]
    if tolerance is not None and booleans["incluye_ruedas_manijas"] is not False:
        errors.append(
            f"Fila {line}, tolerancia_ruedas_manijas_cm: solo tiene sentido cuando "
            "incluye_ruedas_manijas es NO; si la medida ya incluye ruedas y manijas, "
            "o no se sabe, dejar este campo vacío."
        )

    op = {
        "id": text(row.get("id")), "name": name, "type": operator_type,
        "country": country, "region": region, "limit_type": limit_type,
        "notes": text(row.get("notas")), "notes_en": text(row.get("notes_en")),
        "source_url": source, "verified_on": verified.isoformat() if verified else None,
    }
    if limit_type == "dimensional":
        dimensions = {}
        for prefix in ("personal", "cabina"):
            values = [numbers[f"{prefix}_{axis}_cm"] for axis in ("largo", "ancho", "alto")]
            count = sum(value is not None for value in values)
            if count in (1, 2):
                errors.append(f"Fila {line}, {prefix}: completar largo, ancho y alto, o dejar las tres medidas vacías.")
            dimensions[prefix] = dict(zip(("length", "width", "height"), values))
        if not any(all(v is not None for v in dims.values()) for dims in dimensions.values()):
            errors.append(f"Fila {line}: una regla dimensional necesita al menos una terna completa de medidas.")
        op["personal_item"] = {
            "included": booleans["personal_incluido"], "max_cm": dimensions["personal"],
            "max_kg": numbers["personal_kg"],
        }
        op["cabin_bag"] = {
            "included": booleans["cabina_incluida"], "requires_addon": text(row.get("cabina_addon")),
            "max_cm": dimensions["cabina"], "max_kg": numbers["cabina_kg"],
        }
        op["wheels_handles"] = {
            "included_in_measurement": booleans["incluye_ruedas_manijas"],
            "tolerance_cm": tolerance,
        }
        fare = text(row.get("tarifa_referencia"))
        if fare is not None and (len(fare) > MAX_FARE_LEN or "\n" in fare):
            errors.append(f"Fila {line}, tarifa_referencia: usar solo el nombre de la tarifa, "
                          f"en una línea y hasta {MAX_FARE_LEN} caracteres.")
        op["reference_fare"] = fare
        combined = read("peso_combinado_kg", number)
        op["combined_max_kg"] = combined
        for key, label in (("personal_kg", "personal"), ("cabina_kg", "cabina")):
            single = numbers[key]
            if combined is not None and single is not None and single > combined:
                errors.append(f"Fila {line}, {key}: {single} kg supera el peso combinado de "
                              f"{combined} kg; un bulto solo no puede pesar más que los dos juntos.")
        op["excess_fee"] = {
            "currency": currency, "online_from": numbers["tarifa_online_desde"],
            "at_gate_from": numbers["tarifa_en_puerta_desde"],
        }
    elif limit_type in ("descriptivo", "ninguno"):
        op["descriptive_rule"] = {"es": text(row.get("regla_texto_es")), "en": text(row.get("regla_texto_en"))}
        if not any(op["descriptive_rule"].values()):
            errors.append(f"Fila {line}: una regla {limit_type} necesita texto en español o inglés que explique su alcance.")
    if verified:
        age = (today - verified).days
        op["days_since_verified"] = age
        if age > STALE_DAYS:
            warnings.append(f"Fila {line} ({op['id']}): revisión de hace {age} días; conviene reauditar.")
    return op, errors, warnings


def convert_rows(rows, today):
    ids = Counter(text(row.get("id")) for _, row in rows if text(row.get("id")))
    operators, errors, drafts, warnings = [], [], [], []
    for line, row in rows:
        oid = text(row.get("id"))
        if oid is None or not SLUG.fullmatch(oid):
            errors.append(f"Fila {line}, id: usar minúsculas, números y guiones, por ejemplo ryanair.")
            continue
        if ids[oid] > 1:
            errors.append(f"Fila {line}: id duplicado '{oid}'; se bloquean todas sus apariciones.")
            continue
        if text(row.get("verificado_el")) is None:
            drafts.append(f"Fila {line} ({oid}): borrador omitido; falta verificado_el cargado por una persona.")
            continue
        op, row_errors, row_warnings = build_operator(row, line, today)
        errors.extend(row_errors)
        warnings.extend(row_warnings)
        if not row_errors:
            operators.append(op)
    return operators, errors, drafts, warnings


def atomic_write(path, content):
    path = Path(path)
    # La carpeta de destino puede no existir todavía: en un clon limpio del
    # repositorio, public/ no está hasta que algo la crea. Sin esto, correr el
    # conversor recién clonado fallaba al escribir el temporal.
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="\n", dir=path.parent,
                                         prefix=f".{path.name}.", suffix=".tmp", delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(content)
        os.replace(temporary, path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", nargs="?", default="willitboard_operadores.xlsx", help="planilla de entrada")
    parser.add_argument("destination", nargs="?", default="public/operators.json", help="JSON de publicación")
    parser.add_argument("--report", type=Path, help="guardar además un informe de texto UTF-8")
    args = parser.parse_args(argv)
    src, dst = Path(args.source), Path(args.destination)
    paths = [src.resolve(), dst.resolve()] + ([args.report.resolve()] if args.report else [])
    if len(paths) != len(set(paths)):
        parser.error("la planilla, el JSON y el informe deben usar rutas diferentes")
    today = date.today()
    try:
        rows, examples = load_rows(src)
        operators, errors, drafts, warnings = convert_rows(rows, today)
    except (ImportError, OSError, ValueError, KeyError, BadZipFile) as exc:
        errors = [f"No se pudo leer la planilla: {exc}"]
        operators, drafts, warnings, examples = [], [], [], []
    lines = [f"Willitboard — validación {today.isoformat()} — esquema {SCHEMA_VERSION}"]
    lines.append(f"Operadores válidos: {len(operators)}. Borradores omitidos: {len(drafts)}. Ejemplos omitidos: {len(examples)}.")
    lines.append("La URL se valida por su formato; la comprobación de la fuente oficial y la fecha de revisión corresponde a una persona.")
    if examples:
        lines.append("Ejemplos identificados por su nota, en filas: " + ", ".join(map(str, examples)) + ".")
    lines.extend(drafts)
    lines.extend(warnings)
    if errors:
        lines.append(f"ERRORES ({len(errors)}): no se escribe el JSON; cualquier publicación existente se conserva.")
        lines.extend(errors)
    else:
        payload = {
            "schema_version": SCHEMA_VERSION, "generated_on": today.isoformat(),
            "canonical_units": {"length": "cm", "weight": "kg"}, "operators": operators,
        }
        try:
            atomic_write(dst, json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
            if operators:
                lines.append(f"Publicados {len(operators)} operadores en {dst}.")
            else:
                lines.append(f"No hay operadores publicables: se genera {dst} con operators vacío. No es un catálogo verificado.")
        except OSError as exc:
            errors.append(str(exc))
            lines.append(f"ERROR al escribir el JSON: {exc}")
    report = "\n".join(lines) + "\n"
    print(report, end="")
    if args.report:
        try:
            atomic_write(args.report, report)
        except OSError as exc:
            print(f"ERROR al guardar el informe: {exc}", file=sys.stderr)
            return 1
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
