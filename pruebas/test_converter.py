#!/usr/bin/env python3
"""Pruebas automatizadas para convert.py (esquema 1.4).

Reemplaza a pruebas/test_converter_v11.py: ese archivo no estaba disponible
(no se encontró su ubicación) y en cualquier caso no cubría el esquema 1.2.
Sin rutas hardcodeadas: importa convert.py por nombre de módulo, buscándolo
junto a este archivo o en el directorio de trabajo actual. Poné ambos
archivos en la misma carpeta y listo, corra donde corra.

Uso: python -m unittest test_converter.py -v
"""
from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl


def _load_convert_module():
    here = Path(__file__).resolve().parent
    for candidate in (here / "convert.py", Path.cwd() / "convert.py"):
        if candidate.exists():
            spec = importlib.util.spec_from_file_location("convert", candidate)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            return module
    raise FileNotFoundError(
        "No se encontró convert.py junto a este archivo de pruebas ni en el "
        "directorio de trabajo actual. Colocá ambos archivos en la misma carpeta."
    )


convert = _load_convert_module()
HEADERS = convert.REQUIRED_HEADERS

BASE_ROW = {
    "id": "testop", "nombre": "Test Operator", "tipo": "air", "pais": "ES",
    "region": "EU", "limit_type": "dimensional", "personal_incluido": "SI",
    "personal_largo_cm": 40, "personal_ancho_cm": 30, "personal_alto_cm": 20,
    "personal_kg": None, "cabina_incluida": "NO", "cabina_addon": "Priority",
    "cabina_largo_cm": 55, "cabina_ancho_cm": 40, "cabina_alto_cm": 20,
    "cabina_kg": 10, "incluye_ruedas_manijas": None,
    "tolerancia_ruedas_manijas_cm": None, "moneda": None, "tarifa_online_desde": None,
    "tarifa_en_puerta_desde": None, "regla_texto_es": None, "regla_texto_en": None,
    "notas": "fila de prueba", "fuente_url": "https://example.com/baggage",
    "verificado_el": "2026-01-01",
}


def make_workbook(rows, tmp_path, name="planilla.xlsx"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Operadores")
    ws.append(HEADERS)
    ws.append([""] * len(HEADERS))  # fila 2: pistas, el conversor no la usa
    for row in rows:
        ws.append([row.get(h) for h in HEADERS])
    path = tmp_path / name
    wb.save(path)
    return path


class ParsersTests(unittest.TestCase):
    def test_boolean_valores(self):
        self.assertIsNone(convert.boolean(None))
        self.assertIsNone(convert.boolean(""))
        self.assertTrue(convert.boolean("SI"))
        self.assertTrue(convert.boolean("sí"))
        self.assertFalse(convert.boolean("NO"))
        with self.assertRaises(ValueError):
            convert.boolean("tal vez")

    def test_number_positivo(self):
        self.assertEqual(convert.number("40"), 40)
        self.assertEqual(convert.number("40,5"), 40.5)
        with self.assertRaises(ValueError):
            convert.number("0")
        with self.assertRaises(ValueError):
            convert.number("-1")
        with self.assertRaises(ValueError):
            convert.number("abc")

    def test_number_allow_zero(self):
        self.assertEqual(convert.number("0", allow_zero=True), 0)
        with self.assertRaises(ValueError):
            convert.number("-1", allow_zero=True)

    def test_review_date(self):
        today = date(2026, 9, 12)
        self.assertEqual(convert.review_date("2026-09-12", today), today)
        with self.assertRaises(ValueError):
            convert.review_date("2026-09-13", today)
        with self.assertRaises(ValueError):
            convert.review_date("13/09/2026", today)

    def test_source_url(self):
        self.assertEqual(convert.source_url("https://help.ryanair.com/x"),
                          "https://help.ryanair.com/x")
        with self.assertRaises(ValueError):
            convert.source_url("no-es-una-url")
        with self.assertRaises(ValueError):
            convert.source_url(None)


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def _convert(self, rows):
        src = make_workbook(rows, self.tmp)
        rows_loaded, examples = convert.load_rows(src)
        return convert.convert_rows(rows_loaded, date(2026, 9, 12))

    def test_fila_basica_valida(self):
        operators, errors, drafts, warnings = self._convert([dict(BASE_ROW)])
        self.assertEqual(errors, [])
        self.assertEqual(len(operators), 1)
        self.assertEqual(operators[0]["id"], "testop")

    def test_borrador_sin_fecha_se_omite(self):
        row = dict(BASE_ROW, verificado_el=None)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertEqual(operators, [])
        self.assertEqual(errors, [])
        self.assertEqual(len(drafts), 1)

    def test_id_duplicado_se_rechaza(self):
        rows = [dict(BASE_ROW), dict(BASE_ROW, nombre="Otro")]
        operators, errors, drafts, warnings = self._convert(rows)
        self.assertEqual(operators, [])
        self.assertTrue(any("duplicado" in e for e in errors))

    def test_fecha_futura_se_rechaza(self):
        row = dict(BASE_ROW, verificado_el="2026-12-31")
        operators, errors, drafts, warnings = self._convert([row])
        self.assertTrue(any("futuro" in e for e in errors))

    def test_terna_incompleta_se_rechaza(self):
        row = dict(BASE_ROW, personal_ancho_cm=None)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertTrue(any("largo, ancho y alto" in e for e in errors))

    def test_regla_descriptiva_necesita_texto(self):
        row = dict(BASE_ROW, limit_type="descriptivo",
                   personal_largo_cm=None, personal_ancho_cm=None, personal_alto_cm=None,
                   cabina_largo_cm=None, cabina_ancho_cm=None, cabina_alto_cm=None,
                   regla_texto_es=None, regla_texto_en=None)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertTrue(any("necesita texto" in e for e in errors))

    def test_ejemplo_se_omite_por_nota(self):
        row = dict(BASE_ROW, notas="EJEMPLO SIN VERIFICAR - no usar")
        src = make_workbook([row], self.tmp)
        rows_loaded, examples = convert.load_rows(src)
        self.assertEqual(rows_loaded, [])
        self.assertEqual(len(examples), 1)

    def test_revision_vieja_genera_aviso(self):
        row = dict(BASE_ROW, verificado_el="2025-01-01")
        operators, errors, drafts, warnings = self._convert([row])
        self.assertEqual(errors, [])
        self.assertTrue(any("reauditar" in w for w in warnings))


class RuedasManijasTests(unittest.TestCase):
    """Casos del campo nuevo tolerancia_ruedas_manijas_cm, esquema 1.2 (H3)."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def _convert(self, rows):
        src = make_workbook(rows, self.tmp)
        rows_loaded, examples = convert.load_rows(src)
        return convert.convert_rows(rows_loaded, date(2026, 9, 12))

    def test_tolerancia_sin_no_se_rechaza(self):
        row = dict(BASE_ROW, incluye_ruedas_manijas="SI", tolerancia_ruedas_manijas_cm=5)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertEqual(operators, [])
        self.assertTrue(any("tolerancia_ruedas_manijas_cm" in e for e in errors))

    def test_tolerancia_sin_declarar_incluye_se_rechaza(self):
        row = dict(BASE_ROW, incluye_ruedas_manijas=None, tolerancia_ruedas_manijas_cm=5)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertEqual(operators, [])
        self.assertTrue(any("tolerancia_ruedas_manijas_cm" in e for e in errors))

    def test_tolerancia_negativa_se_rechaza(self):
        row = dict(BASE_ROW, incluye_ruedas_manijas="NO", tolerancia_ruedas_manijas_cm=-3)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertEqual(operators, [])
        self.assertTrue(any("mayor o igual a cero" in e for e in errors))

    def test_tolerancia_cero_es_valida(self):
        row = dict(BASE_ROW, incluye_ruedas_manijas="NO", tolerancia_ruedas_manijas_cm=0)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertEqual(errors, [])
        self.assertEqual(operators[0]["wheels_handles"],
                          {"included_in_measurement": False, "tolerance_cm": 0})

    def test_tolerancia_con_no_se_publica(self):
        row = dict(BASE_ROW, incluye_ruedas_manijas="NO", tolerancia_ruedas_manijas_cm=5)
        operators, errors, drafts, warnings = self._convert([row])
        self.assertEqual(errors, [])
        self.assertEqual(operators[0]["wheels_handles"],
                          {"included_in_measurement": False, "tolerance_cm": 5})

    def test_sin_dato_queda_null_null(self):
        operators, errors, drafts, warnings = self._convert([dict(BASE_ROW)])
        self.assertEqual(errors, [])
        self.assertEqual(operators[0]["wheels_handles"],
                          {"included_in_measurement": None, "tolerance_cm": None})


class ArchivoDeSalidaTests(unittest.TestCase):
    """Pruebas end-to-end contra main(), incluida la regla de no romper una
    publicación existente si la carga nueva tiene errores."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_main_no_pisa_json_existente_si_hay_errores(self):
        good = make_workbook([dict(BASE_ROW)], self.tmp, "buena.xlsx")
        dst = self.tmp / "operators.json"
        rc = convert.main([str(good), str(dst)])
        self.assertEqual(rc, 0)
        original_content = dst.read_text(encoding="utf-8")

        bad_row = dict(BASE_ROW, id="otro", incluye_ruedas_manijas="SI",
                       tolerancia_ruedas_manijas_cm=5)
        bad = make_workbook([bad_row], self.tmp, "mala.xlsx")
        rc2 = convert.main([str(bad), str(dst)])
        self.assertEqual(rc2, 1)
        self.assertEqual(dst.read_text(encoding="utf-8"), original_content)

    def test_main_produce_schema_1_4(self):
        good = make_workbook([dict(BASE_ROW)], self.tmp, "buena2.xlsx")
        dst = self.tmp / "operators.json"
        convert.main([str(good), str(dst)])
        payload = json.loads(dst.read_text(encoding="utf-8"))
        self.assertEqual(payload["schema_version"], "1.4")


class TarifaReferenciaTests(unittest.TestCase):
    """Columna opcional desde el esquema 1.4 (24/09/2026)."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def _convert(self, rows, headers):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Operadores")
        ws.append(headers)
        ws.append([""] * len(headers))
        for row in rows:
            ws.append([row.get(h) for h in headers])
        src = self.tmp / "t.xlsx"
        wb.save(src)
        dst = self.tmp / "o.json"
        rc = convert.main([str(src), str(dst)])
        return rc, (json.loads(dst.read_text(encoding="utf-8")) if dst.exists() else None)

    def test_planilla_vieja_sin_columna_sigue_convirtiendo(self):
        rc, payload = self._convert([dict(BASE_ROW)], list(HEADERS))
        self.assertEqual(rc, 0)
        self.assertIsNone(payload["operators"][0]["reference_fare"])

    def test_tarifa_se_publica_tal_cual(self):
        row = dict(BASE_ROW, tarifa_referencia="  Basic Fare ")
        rc, payload = self._convert([row], list(HEADERS) + ["tarifa_referencia"])
        self.assertEqual(rc, 0)
        self.assertEqual(payload["operators"][0]["reference_fare"], "Basic Fare")

    def test_tarifa_vacia_es_none_nunca_texto_inventado(self):
        rc, payload = self._convert([dict(BASE_ROW, tarifa_referencia="")],
                                    list(HEADERS) + ["tarifa_referencia"])
        self.assertIsNone(payload["operators"][0]["reference_fare"])

    def test_tarifa_con_parrafo_se_rechaza(self):
        row = dict(BASE_ROW, tarifa_referencia="Basic\nincluye un bolso")
        rc, _ = self._convert([row], list(HEADERS) + ["tarifa_referencia"])
        self.assertEqual(rc, 1)


class AlcanceSoloCabinaTests(unittest.TestCase):
    """El comparador es solo de cabina (decisión del 12/09/2026, esquema 1.3).

    Estas pruebas son un guardián: si alguien reintroduce el equipaje facturado
    sin discutirlo, fallan y obligan a leer la decisión antes de seguir.
    """

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_el_esquema_no_pide_columnas_de_facturado(self):
        for campo in ("facturado_kg", "facturado_suma_cm"):
            self.assertNotIn(campo, convert.REQUIRED_HEADERS)

    def test_la_salida_no_trae_bloque_checked(self):
        src = make_workbook([dict(BASE_ROW)], self.tmp)
        rows, _ = convert.load_rows(src)
        operators, errors, _, _ = convert.convert_rows(rows, date(2026, 9, 12))
        self.assertEqual(errors, [])
        self.assertNotIn("checked", operators[0])

    def test_columnas_de_facturado_sobrantes_no_rompen_la_carga(self):
        """Una planilla vieja con las columnas retiradas sigue convirtiendo:
        se ignoran en silencio en lugar de bloquear la publicación."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Operadores")
        headers = list(HEADERS) + ["facturado_kg", "facturado_suma_cm"]
        ws.append(headers)
        ws.append([""] * len(headers))
        fila = dict(BASE_ROW)
        ws.append([fila.get(h) for h in HEADERS] + [23, 275])
        src = self.tmp / "planilla_vieja.xlsx"
        wb.save(src)

        rows, _ = convert.load_rows(src)
        operators, errors, _, _ = convert.convert_rows(rows, date(2026, 9, 12))
        self.assertEqual(errors, [])
        self.assertEqual(len(operators), 1)
        self.assertNotIn("checked", operators[0])


if __name__ == "__main__":
    unittest.main(verbosity=2)
